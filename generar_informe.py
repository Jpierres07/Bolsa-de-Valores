from pathlib import Path
import numpy as np,pandas as pd
from docx import Document
from docx.shared import Inches,Pt
from motor_modelos import preparar,entrenamiento,validacion,escenarios,EXP
P=Path(__file__).resolve().parent
D=Document();sec=D.sections[0];D.styles['Normal'].font.size=Pt(10)
def h(t):D.add_heading(t,level=1)
def p(t):D.add_paragraph(t)
def tabla(df):
    t=D.add_table(rows=1,cols=len(df.columns));t.style='Light Shading Accent 1'
    for c,v in zip(t.rows[0].cells,df.columns):c.text=str(v)
    for row in df.itertuples(index=False,name=None):
        for c,v in zip(t.add_row().cells,row):c.text=f'{v:.3f}' if isinstance(v,(float,np.floating)) else str(v)
D.add_heading('Bolsa de Valores y Oportunidades de Inversión',0)
p('Análisis de series temporales y pronóstico de precios de Southern Copper y Buenaventura bajo tres escenarios, 1996–2027.')
p('Versión unificada con la aplicación Streamlit. Datos históricos: julio de 1996–septiembre de 2026; proyección principal: octubre de 2026–marzo de 2027.')
h('1. Planteamiento del problema')
p('¿Qué modelos de series temporales presentan el mejor desempeño en el ajuste y pronóstico de los cierres ajustados mensuales de Southern Copper y Buenaventura, y qué oportunidades y riesgos se identifican bajo tres escenarios para los próximos seis meses?')
h('2. Objetivo general')
p('Analizar y comparar el desempeño de los modelos de series temporales sobre los precios de cierre ajustado mensuales de SCCO y BVN, mediante errores de entrenamiento y de prueba, para evaluar escenarios exploratorios de inversión a seis meses.')
h('3. Justificación teórica y práctica')
p('Teórica: aplicar descomposición clásica, promedios móviles y suavizamiento exponencial a series bursátiles reales, distinguiendo ajuste de pronóstico. Práctica: comparar rendimientos potenciales y riesgos de dos acciones mineras, sin prometer ganancias.')
h('4. Ficha técnica y depuración')
tabla(pd.DataFrame([['Empresas','SCCO (Southern Copper), BVN (Buenaventura)'],['Mercado','NYSE'],['Variable','Cierre ajustado mensual, USD por acción'],['Fuente','CSV históricos del grupo, originados en Yahoo Finance'],['Periodo','Julio 1996 a septiembre 2026'],['Tamaño','363 meses por acción'],['Contexto','Cobre, oro y plata: Banco Mundial Pink Sheet'],['Herramientas','Python, pandas, statsmodels, Streamlit']],columns=['Campo','Descripción']))
p('Se comprueba unicidad de fechas, continuidad mensual, ausencia de faltantes y positividad del cierre ajustado. Los cambios extremos se identifican y revisan, sin eliminarlos automáticamente por ser potencialmente movimientos reales del mercado.')
h('5. Exploración visual y descomposición')
p('Las gráficas históricas y la descomposición se generan en la web desde los mismos CSV. Modelos clásicos: aditivo Y = T + E + R; multiplicativo Y = T × E × R, periodo estacional de 12 meses. Un componente estacional estimado no demuestra por sí solo estacionalidad estable.')
for t in ['SCCO','BVN']:
    for nombre in ['historico','descomp']:
        im=Path('/mnt/data')/f'informe_{t}_{nombre}.png'
        if im.exists():D.add_picture(str(im),width=Inches(6.1))
h('6. Modelos y criterios de evaluación')
p('Se comparan ocho métodos: PMS3, PMS6, PMS12, PMP3 (pesos crecientes 1, 2, 3), SES, Holt, HW_aditivo y HW_multiplicativo. MAE y RMSE se expresan en USD por acción; MAPE en porcentaje. Todos los errores de entrenamiento se calculan sobre los mismos últimos 339 meses, desde el registro 25, aunque los modelos exponenciales se ajustan con la serie completa.')
h('7. Errores de entrenamiento — cálculo único')
series={}
for t in ['SCCO','BVN']:
    x=pd.read_csv(P/'datos'/f'{t}_mensuales.csv',sep=';',encoding='utf-8-sig',parse_dates=['Fecha']).set_index('Fecha')['Adj Close'];series[t]=preparar(x)
    tab,_,fallos=entrenamiento(series[t]);tab.to_csv(P/f'{t}_errores_entrenamiento.csv',index=False,encoding='utf-8-sig')
    D.add_heading(t,level=2);tabla(tab[['Modelo','MAE','RMSE','MAPE (%)','Meses evaluados']]);p(f'Menor RMSE en entrenamiento: {tab.iloc[0]["Modelo"]} ({tab.iloc[0]["RMSE"]:.4f}).')
h('8. Validación temporal de pronósticos')
p('Se reserva octubre de 2025 a septiembre de 2026 como prueba. Los modelos se ajustan exclusivamente con los datos anteriores; se generan 12 pronósticos desde un solo origen temporal. Los métodos de promedio móvil pronostican un valor constante con la última ventana de entrenamiento, sin usar datos futuros.')
seleccion={}
for t in ['SCCO','BVN']:
    tab,_,fallos=validacion(series[t]);tab.to_csv(P/f'{t}_errores_prueba.csv',index=False,encoding='utf-8-sig')
    elegido=next(m for m in tab['Modelo'] if m in EXP);seleccion[t]=elegido
    D.add_heading(t,level=2);tabla(tab[['Modelo','MAE','RMSE','MAPE (%)','Meses evaluados']]);p(f'Modelo exponencial seleccionado para escenarios: {elegido}, por menor RMSE de prueba entre los modelos de suavizamiento.')
h('9. Pronósticos y tres escenarios a seis meses')
p('Escenario central: pronóstico puntual del modelo exponencial seleccionado. Escenarios favorable y pesimista: cuantiles 95 % y 5 % de simulaciones exploratorias con remuestreo de residuos de entrenamiento y escala √h. NO constituyen intervalos de predicción calibrados ni garantizan probabilidades reales de ocurrencia. Se requiere backtesting para validar su cobertura.')
for t in ['SCCO','BVN']:
    e=escenarios(series[t],seleccion[t],h=6);e.to_csv(P/f'{t}_escenarios.csv',encoding='utf-8-sig');ultimo=series[t].iloc[-1]
    D.add_heading(t,level=2)
    resumen=pd.DataFrame([['Pesimista',e.iloc[-1,0],(e.iloc[-1,0]/ultimo-1)*100],['Central',e.iloc[-1,1],(e.iloc[-1,1]/ultimo-1)*100],['Favorable',e.iloc[-1,2],(e.iloc[-1,2]/ultimo-1)*100]],columns=['Escenario','Precio marzo 2027 USD','Rentabilidad potencial %'])
    tabla(resumen)
    p(f'Precio base septiembre de 2026: US$ {ultimo:.2f}. Modelo: {seleccion[t]}.')
h('10. Interpretación gerencial y límites')
p('La acción con mayor rendimiento central proyectado no es necesariamente la mejor inversión. Comparar volatilidad, riesgo de caída, dispersión de escenarios, liquidez y costos. Los cierres ajustados no equivalen al precio exacto de ejecución; una simulación de S/1,000 requiere tipo de cambio, comisiones e impuestos. No se garantiza ganancia alguna.')
h('11. Fuentes y reproducibilidad')
p('Archivos SCCO_mensuales.csv y BVN_mensuales.csv; datos de contexto de Banco Mundial Pink Sheet. Fuente bursátil de origen: Yahoo Finance. El archivo motor_modelos.py contiene las funciones compartidas por el Word y Streamlit. Los archivos CSV de entrenamiento, prueba y escenarios se generan con generar_informe.py.')
D.save(P/'Informe_SCCO_BVN_unificado.docx')
print('Informe unificado guardado')
